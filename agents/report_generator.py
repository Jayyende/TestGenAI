from docx import Document


def create_report(functions, test_cases, bug_report, filepath):

    doc = Document()

    doc.add_heading(
        'TestGenAI Software Testing Report',
        level=1
    )

    doc.add_heading(
        'Functions Found',
        level=2
    )

    for func in functions:
        doc.add_paragraph(
            f"• {func}()"
        )

    doc.add_heading(
        'Generated Test Cases',
        level=2
    )

    for func, cases in test_cases.items():

        doc.add_paragraph(
            f"{func}()",
            style='Heading 3'
        )

        for case in cases:
            doc.add_paragraph(
                case
            )

    doc.add_heading(
        'Bug Analysis',
        level=2
    )

    doc.add_paragraph(
        f"Risk Level: {bug_report['risk_level']}"
    )

    for issue in bug_report["issues"]:
        doc.add_paragraph(issue)

    doc.add_heading(
        'Conclusion',
        level=2
    )

    doc.add_paragraph(
        "Analysis completed successfully using TestGenAI."
    )

    doc.save(filepath)